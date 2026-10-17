import json
import zipfile

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START

from manuscript_placeholders.cli import main
from manuscript_placeholders.collector import collect
from manuscript_placeholders.review import extra_docx_blocks


def fixture(path):
    d = Document()
    d.add_paragraph("[RESEARCH body]")
    section = d.sections[0]
    section.header.paragraphs[0].text = "[RESEARCH header]"
    section.footer.paragraphs[0].text = "[RESEARCH footer]"
    section.first_page_header.paragraphs[0].text = "[RESEARCH first]"
    section.even_page_footer.paragraphs[0].text = "[RESEARCH even]"
    section.header.add_table(rows=1, cols=1, width=100).cell(
        0, 0
    ).text = "[RESEARCH table]"
    d.add_section(WD_SECTION_START.NEW_PAGE)
    d.save(path)


def test_opt_in_shared_parts_locations_and_source(tmp_path):
    path = tmp_path / "source.docx"
    fixture(path)
    before = path.read_bytes()
    assert len(collect(path)) == 1
    findings = collect(path, docx_headers=True, docx_footers=True)
    assert {x.marker for x in findings} == {
        "[RESEARCH body]",
        "[RESEARCH header]",
        "[RESEARCH footer]",
        "[RESEARCH first]",
        "[RESEARCH even]",
    }
    assert len({x.fingerprint for x in findings}) == len(findings)
    assert any("word/header" in x.location for x in findings)
    assert (
        len(collect(path, docx_headers=True, docx_footers=True, docx_tables=True)) == 6
    )
    assert path.read_bytes() == before


def test_cli_coverage_flags(tmp_path, capsys):
    path = tmp_path / "source.docx"
    fixture(path)
    assert main([str(path), "--docx-headers", "--format", "json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["review"]["docx_structures"]["headers"] == "included"
    assert "skipped" in data["review"]["docx_structures"]["footers"]
    assert not any("footer" in x["location"] for x in data["placeholders"])


@pytest.mark.parametrize(
    "mutation", ["external", "unsafe", "missing", "bad-root", "malformed"]
)
def test_invalid_header_parts(tmp_path, mutation):
    source = tmp_path / "source.docx"
    fixture(source)
    with zipfile.ZipFile(source) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    name = "word/header1.xml"
    if mutation in ("external", "unsafe"):
        rel = parts["word/_rels/document.xml.rels"].decode()
        if mutation == "external":
            rel = rel.replace(
                'Target="header1.xml"',
                'Target="https://invalid.example/x" TargetMode="External"',
            )
        else:
            rel = rel.replace('Target="header1.xml"', 'Target="../../escape.xml"')
        parts["word/_rels/document.xml.rels"] = rel.encode()
    elif mutation == "missing":
        del parts[name]
    elif mutation == "bad-root":
        parts[name] = b"<wrong/>"
    else:
        parts[name] = b"<broken"
    with zipfile.ZipFile(source, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    with pytest.raises(
        (
            ValueError,
            __import__("xml.etree.ElementTree", fromlist=["ParseError"]).ParseError,
        )
    ):
        extra_docx_blocks(
            source, tables=False, comments=False, footnotes=False, headers=True
        )
