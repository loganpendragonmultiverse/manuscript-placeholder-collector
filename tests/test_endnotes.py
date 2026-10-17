import importlib
import json
import os
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document

from manuscript_placeholders import output
from manuscript_placeholders.cli import main
from manuscript_placeholders.collector import collect

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"
NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
NOTE = (
    f"<w:endnotes {NS}>"
    '<w:endnote w:id="-1" w:type="separator"><w:p><w:r><w:t>[CHECK separator]</w:t></w:r></w:p></w:endnote>'
    '<w:endnote w:id="0" w:type="continuationSeparator"><w:p/></w:endnote>'
    '<w:endnote w:id="7"><w:p/>'
    '<w:p><w:r><w:t xml:space="preserve">  [RE</w:t></w:r><w:r><w:t>SEARCH note]</w:t></w:r></w:p>'
    "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>[VERIFY table]</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
    "<w:p><w:r><w:t>[CHECK after table]</w:t></w:r></w:p></w:endnote>"
    '<w:endnote w:id="22"><w:p><w:r><w:t>[RESEARCH note]</w:t></w:r></w:p></w:endnote>'
    "</w:endnotes>"
)


def package(path, note=NOTE, target="endnotes.xml", references=(7, 7, 22)):
    d = Document()
    d.add_paragraph("[RESEARCH body]")
    d.save(path)
    with ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    document = ET.fromstring(parts["word/document.xml"])
    body = document.find(W + "body")
    p = ET.Element(W + "p")
    for identifier in references:
        run = ET.SubElement(p, W + "r")
        ET.SubElement(run, W + "endnoteReference", {W + "id": str(identifier)})
    body.insert(1, p)
    parts["word/document.xml"] = ET.tostring(document)
    rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
    ET.SubElement(
        rels,
        REL + "Relationship",
        {
            "Id": "rIdEndnotes",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes",
            "Target": target,
        },
    )
    parts["word/_rels/document.xml.rels"] = ET.tostring(rels)
    name = "word/" + target if not target.startswith("/") else target.lstrip("/")
    types = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(
        types,
        CT + "Override",
        {
            "PartName": "/" + name,
            "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
        },
    )
    parts["[Content_Types].xml"] = ET.tostring(types)
    parts[name] = note.encode()
    save(path, parts)
    return parts


def save(path, parts):
    with ZipFile(path, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)


def test_endnotes_opt_in_ids_positions_offsets_and_source(tmp_path):
    path = tmp_path / "notes.docx"
    package(path)
    before = path.read_bytes()
    assert len(collect(path)) == 1
    found = collect(path, docx_endnotes=True)
    assert len(found) == 4
    selected = [x for x in found if "/endnote[" in x.location]
    assert [x.location for x in selected] == [
        "word/endnotes.xml/endnote[7]/paragraph[2]",
        "word/endnotes.xml/endnote[7]/paragraph[4]",
        "word/endnotes.xml/endnote[22]/paragraph[1]",
    ]
    assert selected[0].start == 2
    assert selected[0].end == 17
    assert len({x.fingerprint for x in found}) == 4
    expanded = collect(path, docx_endnotes=True, docx_tables=True)
    assert len(expanded) == 5
    assert next(x.location for x in expanded if x.marker == "[VERIFY table]") == (
        "word/endnotes.xml/endnote[7]/paragraph[3]"
    )
    assert {x.fingerprint for x in found} <= {x.fingerprint for x in expanded}
    assert collect(path, docx_endnotes=True) == found
    assert path.read_bytes() == before


@pytest.mark.parametrize("target", ["notes/custom.xml", "/word/endnotes.xml"])
def test_declared_local_part_path(tmp_path, target):
    path = tmp_path / "notes.docx"
    package(path, target=target)
    found = collect(path, docx_endnotes=True)
    assert (
        any(x.location.startswith("word/" + target + "/") for x in found)
        if not target.startswith("/")
        else any(x.location.startswith("word/endnotes.xml/") for x in found)
    )


def test_no_endnotes_is_valid_and_unrelated_orphan_part_is_not_scanned(tmp_path):
    path = tmp_path / "plain.docx"
    d = Document()
    d.add_paragraph("[RESEARCH body]")
    d.save(path)
    with ZipFile(path, "a") as z:
        z.writestr("word/endnotes.xml", NOTE)
    assert collect(path, docx_endnotes=True) == collect(path)


def test_footnote_and_endnote_ids_remain_distinct(tmp_path):
    path = tmp_path / "notes.docx"
    parts = package(path)
    parts["word/footnotes.xml"] = (
        f'<w:footnotes {NS}><w:footnote w:id="7"><w:p><w:r><w:t>'
        "[RESEARCH note]</w:t></w:r></w:p></w:footnote></w:footnotes>"
    ).encode()
    save(path, parts)
    from xml.etree import ElementTree as ET

    relpath = "word/_rels/document.xml.rels"
    rels = ET.fromstring(parts[relpath])
    ET.SubElement(
        rels,
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
        Id="footnotes",
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
        Target="footnotes.xml",
    )
    parts[relpath] = ET.tostring(rels)
    save(path, parts)
    found = collect(path, docx_endnotes=True, docx_footnotes=True)
    assert any("/footnote[7]/" in x.location for x in found)
    assert any("/endnote[7]/" in x.location for x in found)
    assert len({x.fingerprint for x in found}) == len(found)


def test_text_boxes_remain_excluded(tmp_path):
    path = tmp_path / "notes.docx"
    note = NOTE.replace(
        '<w:endnote w:id="7"><w:p/>',
        '<w:endnote w:id="7"><w:p><w:r><w:drawing><w:txbxContent><w:p><w:r><w:t>[RESEARCH textbox]</w:t></w:r></w:p></w:txbxContent></w:drawing></w:r></w:p>',
    )
    package(path, note=note)
    found = collect(path, docx_endnotes=True, docx_tables=True)
    assert "[RESEARCH textbox]" not in {x.marker for x in found}


@pytest.mark.parametrize("format_name", ["json", "html", "markdown", "text"])
def test_cli_reports_endnotes_and_ledger(tmp_path, capsys, format_name):
    source = tmp_path / "notes.docx"
    package(source)
    before = source.read_bytes()
    report, ledger = tmp_path / "report", tmp_path / "ledger.json"
    assert (
        main(
            [
                str(source),
                "--docx-endnotes",
                "--format",
                format_name,
                "--output",
                str(report),
                "--ledger-output",
                str(ledger),
            ]
        )
        == 0
    )
    capsys.readouterr()
    assert "endnotes" in report.read_text()
    assert "endnote[7]" in report.read_text()
    assert len(json.loads(ledger.read_text())["items"]) == 4
    if format_name == "json":
        data = json.loads(report.read_text())
        assert data["review"]["docx_structures"]["endnotes"] == "included"
    assert source.read_bytes() == before


@pytest.mark.parametrize(
    "mutation",
    [
        "malformed",
        "wrong-root",
        "missing-part",
        "external",
        "unsafe",
        "duplicate-id",
        "missing-id",
        "invalid-id",
        "unknown-type",
        "missing-reference",
        "no-relationship",
        "multiple-relationships",
    ],
)
def test_bad_endnotes_fail_before_output_and_preserve_source(
    tmp_path, capsys, mutation
):
    source = tmp_path / "bad.docx"
    parts = package(source)
    if mutation == "malformed":
        parts["word/endnotes.xml"] = b"<broken"
    elif mutation == "wrong-root":
        parts["word/endnotes.xml"] = b"<wrong/>"
    elif mutation == "missing-part":
        del parts["word/endnotes.xml"]
    elif mutation in {
        "external",
        "unsafe",
        "no-relationship",
        "multiple-relationships",
    }:
        rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
        rel = next(x for x in rels if x.get("Id") == "rIdEndnotes")
        if mutation == "external":
            rel.set("TargetMode", "External")
            rel.set("Target", "https://invalid.example/endnotes.xml")
        elif mutation == "unsafe":
            rel.set("Target", "../../endnotes.xml")
        elif mutation == "no-relationship":
            rels.remove(rel)
        else:
            ET.SubElement(
                rels, REL + "Relationship", dict(rel.attrib, Id="rIdOtherNotes")
            )
        parts["word/_rels/document.xml.rels"] = ET.tostring(rels)
    else:
        raw = parts["word/endnotes.xml"].decode()
        if mutation == "duplicate-id":
            raw = raw.replace('w:id="22"', 'w:id="007"')
        elif mutation == "missing-id":
            raw = raw.replace('w:id="7"', "")
        elif mutation == "invalid-id":
            raw = raw.replace('w:id="7"', 'w:id="seven"')
        elif mutation == "unknown-type":
            raw = raw.replace('w:id="7"', 'w:id="7" w:type="unexpected"')
        else:
            raw = raw.replace('w:id="22"', 'w:id="23"')
        parts["word/endnotes.xml"] = raw.encode()
    save(source, parts)
    before = source.read_bytes()
    report, ledger = tmp_path / "report.json", tmp_path / "ledger.json"
    assert (
        main(
            [
                str(source),
                "--docx-endnotes",
                "--format",
                "json",
                "--output",
                str(report),
                "--ledger-output",
                str(ledger),
            ]
        )
        == 2
    )
    assert "Traceback" not in capsys.readouterr().err
    assert not report.exists() and not ledger.exists()
    assert source.read_bytes() == before


def test_endnotes_disabled_and_enabled_coverage(tmp_path, capsys):
    source = tmp_path / "notes.docx"
    package(source)
    assert main([str(source), "--format", "json"]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert "skipped" in payload["review"]["docx_structures"]["endnotes"]
    assert payload["count"] == 1


def test_endnotes_partial_output_failure_and_retry(tmp_path, capsys, monkeypatch):
    source = tmp_path / "notes.docx"
    package(source)
    before = source.read_bytes()
    report, ledger = tmp_path / "report.json", tmp_path / "ledger.json"
    args = [
        str(source),
        "--docx-endnotes",
        "--format",
        "json",
        "--output",
        str(report),
        "--ledger-output",
        str(ledger),
    ]
    original = output._write_final

    def fail(item):
        if item.path == ledger:
            item.written = os.write(item.fd, item.data[:2])
            raise OSError("synthetic endnote output failure")
        return original(item)

    with monkeypatch.context() as patch:
        patch.setattr(output, "_write_final", fail)
        assert main(args) == 2
    capsys.readouterr()
    assert not report.exists() and not ledger.exists()
    assert main(args) == 0
    assert json.loads(report.read_text())["count"] == 4
    assert source.read_bytes() == before


@pytest.mark.parametrize("blocked", ["report", "ledger"])
@pytest.mark.skipif(
    os.name != "posix" or getattr(os, "geteuid", lambda: 0)() == 0,
    reason="requires non-admin POSIX permissions",
)
def test_endnotes_real_destination_failure_and_retry(tmp_path, capsys, blocked):
    source = tmp_path / "notes.docx"
    package(source)
    before = source.read_bytes()
    readonly = tmp_path / "readonly"
    readonly.mkdir()
    readonly.chmod(0o555)
    report = (readonly if blocked == "report" else tmp_path) / "report.json"
    ledger = (readonly if blocked == "ledger" else tmp_path) / "ledger.json"
    args = [
        str(source),
        "--docx-endnotes",
        "--format",
        "json",
        "--output",
        str(report),
        "--ledger-output",
        str(ledger),
    ]
    try:
        assert main(args) == 2
        capsys.readouterr()
        assert not report.exists() and not ledger.exists()
    finally:
        readonly.chmod(0o755)
    assert main(args) == 0
    assert json.loads(report.read_text())["count"] == 4
    assert source.read_bytes() == before


def test_endnotes_stdout_failure_leaves_no_ledger(tmp_path, monkeypatch, capsys):
    source = tmp_path / "notes.docx"
    package(source)
    ledger = tmp_path / "ledger.json"

    class Broken:
        def write(self, text):
            raise BrokenPipeError("closed output")

        def flush(self):
            raise BrokenPipeError("closed output")

    module = importlib.import_module("manuscript_placeholders.cli")
    with monkeypatch.context() as patch:
        patch.setattr(module.sys, "stdout", Broken())
        assert (
            main([str(source), "--docx-endnotes", "--ledger-output", str(ledger)]) == 2
        )
    capsys.readouterr()
    assert not ledger.exists()
