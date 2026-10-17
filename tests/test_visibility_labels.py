import json

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from manuscript_placeholders.cli import main
from manuscript_placeholders.collector import collect


@pytest.mark.parametrize("value", ["0", "false", "off", "1", "true", "on", None])
def test_explicit_visibility_values(tmp_path, capsys, value):
    source = tmp_path / "draft.docx"
    document = Document()
    run = document.add_paragraph().add_run("[TODO retain marker]")
    vanish = OxmlElement("w:vanish")
    if value is not None:
        vanish.set(qn("w:val"), value)
    run._r.get_or_add_rPr().append(vanish)
    document.save(source)
    before = source.read_bytes()
    assert main([str(source), "--format", "json"]) == 0
    finding = json.loads(capsys.readouterr().out)["placeholders"][0]
    assert finding["marker"] == "[TODO retain marker]"
    assert ("paragraph contains hidden text" in finding["section"]) == (
        value not in {"0", "false", "off"}
    )
    assert source.read_bytes() == before


def test_mixed_paragraph_label_describes_context(tmp_path):
    source = tmp_path / "mixed.docx"
    document = Document()
    paragraph = document.add_paragraph("[TODO visible] ")
    paragraph.add_run("hidden context only").font.hidden = True
    document.save(source)
    before = source.read_bytes()
    findings = collect(source)
    assert [item.marker for item in findings] == ["[TODO visible]"]
    assert findings[0].section == "Document [paragraph contains hidden text]"
    assert source.read_bytes() == before
