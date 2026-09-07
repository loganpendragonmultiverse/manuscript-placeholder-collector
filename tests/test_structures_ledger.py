import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from manuscript_placeholders.cli import main
from manuscript_placeholders.collector import collect
from manuscript_placeholders.review import review_findings


def test_opt_in_docx_tables_comments_footnotes_have_locations(tmp_path: Path) -> None:
    source = tmp_path / "draft.docx"
    document = Document()
    document.add_paragraph("Body [RESEARCH body]")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "[VERIFY table]"
    document.save(str(source))
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    with ZipFile(source, "a") as archive:
        archive.writestr(
            "word/comments.xml",
            f'<w:comments {ns}><w:comment w:id="9"><w:p><w:r><w:t>[CHECK comment]</w:t></w:r></w:p></w:comment></w:comments>',
        )
        archive.writestr(
            "word/footnotes.xml",
            f'<w:footnotes {ns}><w:footnote w:id="3"><w:p><w:r><w:t>[SOURCE footnote]</w:t></w:r></w:p></w:footnote></w:footnotes>',
        )
    original = source.read_bytes()
    assert len(collect(source)) == 1
    found = collect(source, docx_tables=True, docx_comments=True, docx_footnotes=True)
    assert len(found) == 4
    assert any("table[1]/row[1]/cell[1]" in item.location for item in found)
    assert any("comment[9]" in item.location for item in found)
    assert any("footnote[3]" in item.location for item in found)
    assert all(item.end > item.start and item.fingerprint for item in found)
    report, ledger = tmp_path / "report.json", tmp_path / "ledger.json"
    assert (
        main(
            [
                str(source),
                "--docx-tables",
                "--docx-comments",
                "--docx-footnotes",
                "--format",
                "json",
                "--output",
                str(report),
                "--ledger-output",
                str(ledger),
            ]
        )
        == 0
    )
    assert len(json.loads(ledger.read_text(encoding="utf-8"))["items"]) == 4
    assert source.read_bytes() == original


def test_resolution_and_comparison_preserve_author_assertions(tmp_path: Path) -> None:
    source = tmp_path / "draft.md"
    source.write_text("  [RESEARCH original]", encoding="utf-8")
    item = collect(source)[0]
    assert item.start == 2
    ledger = {
        "version": 1,
        "items": {
            item.fingerprint: {"status": "resolved", "note": "Private review note"}
        },
    }
    review = review_findings([], ledger, {"placeholders": [item.to_dict()]})
    assert review["ledger"] == ledger
    assert review["comparison"]["no_longer_observed"] == [item.fingerprint]
    assert review["not_currently_observed"] == [item.fingerprint]
    for bad in ([], {"items": []}, {"items": {"x": {"status": "made-up"}}}):
        with pytest.raises((TypeError, ValueError)):
            review_findings([], bad)
    with pytest.raises((TypeError, ValueError)):
        review_findings([], {}, {"placeholders": [None]})
