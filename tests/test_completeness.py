import json
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from manuscript_placeholders.cli import main
from manuscript_placeholders.collector import collect
from manuscript_placeholders.docx_reader import W


def docx(path: Path, content: str, rels: str = "") -> None:
    document = Document()
    document.save(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["word/document.xml"] = (
        f'<w:document xmlns:w="{W[1:-1]}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><w:body>{content}</w:body></w:document>'
    ).encode()
    if rels:
        root = ET.fromstring(parts["word/_rels/document.xml.rels"])
        ET.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            Id="custom",
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
            + rels,
            Target="unused.xml",
        )
        parts["word/_rels/document.xml.rels"] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as z:
        for n, data in parts.items():
            z.writestr(n, data)


def paragraph(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def test_sdt_customxml_textbox_no_double_scan(tmp_path):
    path = tmp_path / "draft.docx"
    docx(
        path,
        "<w:sdt><w:sdtContent>"
        + paragraph("[TODO control]")
        + "</w:sdtContent></w:sdt><w:customXml>"
        + paragraph("[RESEARCH custom]")
        + "</w:customXml><w:p><w:r><w:t>[NAME outside]</w:t><w:drawing><w:txbxContent>"
        + paragraph("[PLACE inside]")
        + "</w:txbxContent></w:drawing></w:r></w:p>",
    )
    original = path.read_bytes()
    assert len(collect(path)) == 3
    found = collect(path, docx_textboxes=True)
    assert len(found) == 4
    assert len({f.fingerprint for f in found}) == 4
    assert path.read_bytes() == original


def test_alternate_branch_once(tmp_path):
    path = tmp_path / "draft.docx"
    docx(
        path,
        '<mc:AlternateContent><mc:Choice Requires="w">'
        + paragraph("[TODO choice]")
        + "</mc:Choice><mc:Fallback>"
        + paragraph("[TODO duplicate]")
        + "</mc:Fallback></mc:AlternateContent>",
    )
    assert [x.marker for x in collect(path)] == ["[TODO choice]"]


def test_revisions_tabs_breaks_fields_hidden(tmp_path):
    path = tmp_path / "draft.docx"
    docx(
        path,
        "<w:del>"
        + paragraph("[TODO deleted paragraph]")
        + "</w:del><w:ins>"
        + paragraph("[TODO inserted]")
        + '</w:ins><w:p><w:r><w:rPr><w:vanish/></w:rPr><w:t>[RESEARCH a</w:t><w:tab/><w:t>b</w:t><w:br/><w:t>c]</w:t></w:r><w:del><w:r><w:delText>[NAME old]</w:delText></w:r></w:del><w:r><w:fldChar w:fldCharType="begin"/><w:instrText>[TODO instruction]</w:instrText><w:fldChar w:fldCharType="separate"/><w:t>[NAME result]</w:t><w:fldChar w:fldCharType="end"/></w:r></w:p>',
    )
    current = collect(path)
    assert {x.marker for x in current} == {
        "[TODO inserted]",
        "[RESEARCH a\tb\nc]",
        "[NAME result]",
    }
    assert any("hidden" in x.section for x in current)
    assert len(collect(path, revisions="include-revisions")) == 5


def test_coverage_absent_unsupported_and_non_docx_flags(tmp_path, capsys):
    path = tmp_path / "draft.docx"
    docx(path, paragraph("[TODO x]"), "aFChunk")
    assert main([str(path), "--format", "json", "--docx-comments"]) == 0
    payload = json.loads(capsys.readouterr().out)
    coverage = payload["review"]["coverage"]
    assert coverage["stories"]["comments"]["status"] == "absent"
    assert coverage["unsupported"][0]["type"] == "aFChunk"
    assert coverage["complete_selected_scope"] is False
    text = tmp_path / "draft.txt"
    text.write_text("[TODO x]")
    assert main([str(text), "--docx-comments"]) == 2
    assert "only to .docx" in capsys.readouterr().err


def test_document_identity_and_status_totals(tmp_path, capsys):
    path = tmp_path / "one.txt"
    path.write_text("[TODO x]")
    report = tmp_path / "report.json"
    ledger = tmp_path / "ledger.json"
    assert (
        main(
            [
                str(path),
                "--format",
                "json",
                "--output",
                str(report),
                "--ledger-output",
                str(ledger),
            ]
        )
        == 0
    )
    payload = json.loads(ledger.read_text())
    assert payload["version"] == 2
    item = next(iter(payload["items"].values()))
    item["status"] = "resolved"
    ledger.write_text(json.dumps(payload))
    assert main([str(path), "--format", "json", "--ledger", str(ledger)]) == 0
    data = json.loads(capsys.readouterr().out.split("\n", 1)[-1])
    assert data["review"]["status_totals"] == {"open": 0, "resolved": 1, "ignored": 0}
    other = tmp_path / "two.txt"
    other.write_text("[TODO x]")
    assert main([str(other), "--ledger", str(ledger)]) == 2
    assert "another document" in capsys.readouterr().err
    assert collect(path)[0].fingerprint != collect(other)[0].fingerprint
    assert main([str(other), "--ledger", str(ledger), "--allow-cross-document"]) == 0


def test_changed_scope_removals_are_unverified(tmp_path, capsys):
    path = tmp_path / "draft.docx"
    docx(
        path,
        paragraph("[TODO body]")
        + "<w:tbl><w:tr><w:tc>"
        + paragraph("[TODO table]")
        + "</w:tc></w:tr></w:tbl>",
    )
    report = tmp_path / "baseline.json"
    assert (
        main([str(path), "--docx-tables", "--format", "json", "--output", str(report)])
        == 0
    )
    capsys.readouterr()
    assert main([str(path), "--baseline", str(report), "--format", "json"]) == 0
    review = json.loads(capsys.readouterr().out)["review"]
    assert len(review["unverified_removals"]) == 1
    assert review["comparison"]["no_longer_observed"] == []


def test_multiline_positions_and_builtin_overlap(tmp_path):
    path = tmp_path / "draft.txt"
    path.write_text("[TODO first\nsecond]\nOther [TODO third]")
    assert len(collect(path, multiline=True)) == 2
    finding = collect(path, multiline=True)[0]
    assert (
        finding.start_line,
        finding.start_column,
        finding.end_line,
        finding.end_column,
    ) == (1, 1, 2, 8)


@pytest.mark.parametrize("pattern", [".*", "(?=x)", "(a+)+$"])
def test_custom_patterns_bounded_no_partial_outputs(tmp_path, capsys, pattern):
    path = tmp_path / "draft.txt"
    path.write_text("x" if pattern != "(a+)+$" else "a" * 50000 + "!")
    report = tmp_path / "report.json"
    ledger = tmp_path / "ledger.json"
    assert (
        main(
            [
                str(path),
                "--pattern",
                pattern,
                "--output",
                str(report),
                "--ledger-output",
                str(ledger),
            ]
        )
        == 2
    )
    assert not report.exists() and not ledger.exists()
    assert "pattern" in capsys.readouterr().err.lower()


def test_duplicate_package_rejected_before_parsing(tmp_path, capsys):
    path = tmp_path / "bad.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", "x")
        with pytest.warns(UserWarning):
            z.writestr("word/document.xml", "y")
    assert main([str(path)]) == 2
    assert "duplicate" in capsys.readouterr().err


def test_decoder_and_converter_timeout_leave_no_report(tmp_path, capsys, monkeypatch):
    path = tmp_path / "bad.txt"
    path.write_bytes(b"\xff")
    output = tmp_path / "report"
    assert main([str(path), "--output", str(output)]) == 2
    assert not output.exists()
    path = tmp_path / "draft.doc"
    path.write_bytes(b"legacy")
    monkeypatch.setattr(
        "manuscript_placeholders.collector.shutil.which", lambda _: "/converter"
    )

    def timeout(*args, **kwargs):
        assert kwargs["timeout"] == 30
        raise subprocess.TimeoutExpired("converter", 30)

    monkeypatch.setattr("manuscript_placeholders.collector.subprocess.run", timeout)
    assert main([str(path), "--output", str(output)]) == 2
    assert not output.exists()
    assert "timed out" in capsys.readouterr().err


def test_alternate_multiple_choices(tmp_path):
    path = tmp_path / "draft.docx"
    docx(
        path,
        '<mc:AlternateContent><mc:Choice Requires="w">'
        + paragraph("[TODO selected]")
        + '</mc:Choice><mc:Choice Requires="w">'
        + paragraph("[TODO unselected]")
        + "</mc:Choice><mc:Fallback>"
        + paragraph("[TODO fallback]")
        + "</mc:Fallback></mc:AlternateContent>",
    )
    assert [x.marker for x in collect(path)] == ["[TODO selected]"]


@pytest.mark.parametrize(
    "which,payload",
    [
        ("baseline", [1]),
        (
            "ledger",
            {"version": 2, "provenance": ["bad"], "items": {"x": {"status": "open"}}},
        ),
    ],
)
def test_malformed_review_objects_actionable(tmp_path, capsys, which, payload):
    path = tmp_path / "draft.txt"
    path.write_text("[TODO x]")
    prior = tmp_path / "prior.json"
    prior.write_text(json.dumps(payload))
    output = tmp_path / "report"
    assert main([str(path), "--" + which, str(prior), "--output", str(output)]) == 2
    assert not output.exists()
    assert "Traceback" not in capsys.readouterr().err


def test_custom_policy_removal_unverified(tmp_path, capsys):
    path = tmp_path / "draft.txt"
    path.write_text("CUSTOMMARKER")
    prior = tmp_path / "prior.json"
    assert (
        main(
            [
                str(path),
                "--pattern",
                "CUSTOMMARKER",
                "--format",
                "json",
                "--output",
                str(prior),
            ]
        )
        == 0
    )
    capsys.readouterr()
    assert main([str(path), "--baseline", str(prior), "--format", "json"]) == 0
    review = json.loads(capsys.readouterr().out)["review"]
    assert len(review["unverified_removals"]) == 1
    assert not review["comparison"]["no_longer_observed"]


def test_html_actual_unsupported_coverage(tmp_path):
    path = tmp_path / "draft.docx"
    docx(path, paragraph("No markers"), "aFChunk")
    output = tmp_path / "report.html"
    assert main([str(path), "--format", "html", "--output", str(output)]) == 0
    content = output.read_text()
    assert "unsupported" in content and "Actual scan coverage" in content


def test_selected_header_container_coverage(tmp_path):
    path = tmp_path / "header.docx"
    document = Document()
    header = document.sections[0].header
    header.add_table(rows=1, cols=1, width=__import__("docx").shared.Inches(2)).cell(
        0, 0
    ).text = "[TODO header table]"
    box = ET.fromstring(
        '<w:r xmlns:w="'
        + W[1:-1]
        + '"><w:drawing><w:txbxContent>'
        + paragraph("[TODO header box]")
        + "</w:txbxContent></w:drawing></w:r>"
    )
    header.paragraphs[0]._p.append(__import__("docx").oxml.parse_xml(ET.tostring(box)))
    document.save(path)
    coverage = {}
    found = collect(
        path,
        docx_headers=True,
        docx_tables=True,
        docx_textboxes=True,
        coverage=coverage,
    )
    assert len(found) == 2
    for kind in ("tables", "textboxes"):
        assert coverage["stories"][kind]["status"] == "scanned"
        assert coverage["stories"][kind]["count"] == 1
    coverage = {}
    assert not collect(path, docx_headers=True, coverage=coverage)
    assert coverage["stories"]["tables"]["status"] == "skipped"
    assert coverage["stories"]["textboxes"]["status"] == "skipped"


def test_custom_exact_builtin_overlap_kept(tmp_path):
    path = tmp_path / "draft.txt"
    path.write_text("[TODO exact]")
    found = collect(path, (r"\[TODO exact\]",))
    assert len(found) == 2
    assert {item.category for item in found} == {"Custom", "Revision notes"}
    assert len({item.fingerprint for item in found}) == 2


def test_multiline_preserves_markdown_sections(tmp_path):
    path = tmp_path / "draft.md"
    path.write_text("# First\n[TODO first]\n# Second\n[TODO second]\n")
    assert [item.section for item in collect(path, multiline=True)] == [
        "First",
        "Second",
    ]


def test_multiline_short_marker_after_large_prefix(tmp_path):
    path = tmp_path / "draft.txt"
    path.write_text("a" * 8180 + "\n[RESEARCH x\nmore]")
    found = collect(path, multiline=True)
    assert len(found) == 1
    assert found[0].marker == "[RESEARCH x\nmore]"
    assert (
        found[0].start_line,
        found[0].start_column,
        found[0].end_line,
        found[0].end_column,
    ) == (2, 1, 3, 6)


def test_oversized_marker_actionable_no_report(tmp_path, capsys):
    path = tmp_path / "draft.txt"
    path.write_text("[RESEARCH " + "x" * 9000 + "]")
    report = tmp_path / "report"
    assert main([str(path), "--multiline", "--output", str(report)]) == 2
    assert not report.exists()
    assert "8192" in capsys.readouterr().err


@pytest.mark.parametrize("xml", ["<wrong/>", '<w:document xmlns:w="' + W[1:-1] + '"/>'])
def test_invalid_main_docx_root_rejected(tmp_path, capsys, xml):
    path = tmp_path / "draft.docx"
    docx(path, paragraph("text"))
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["word/document.xml"] = xml.encode()
    with zipfile.ZipFile(path, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)
    assert main([str(path)]) == 2
    assert "document body" in capsys.readouterr().err
