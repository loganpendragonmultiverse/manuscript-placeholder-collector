import subprocess
from pathlib import Path

from docx import Document

from manuscript_placeholders import collector


def test_libreoffice_uses_disposable_profile_and_preserves_source(
    tmp_path, monkeypatch
):
    source = tmp_path / "draft.doc"
    source.write_bytes(b"synthetic legacy document")
    commands = []
    monkeypatch.setattr(
        collector.shutil,
        "which",
        lambda tool: "/usr/bin/libreoffice" if tool == "soffice" else None,
    )

    def convert(args, **kwargs):
        commands.append(args)
        target = Path(args[args.index("--outdir") + 1])
        assert "-env:UserInstallation=" + (target / "profile").as_uri() in args
        document = Document()
        document.add_paragraph("TODO: lighthouse")
        document.save(target / "draft.docx")
        return subprocess.CompletedProcess(args, 0, "", "")

    monkeypatch.setattr(collector.subprocess, "run", convert)
    result = collector.collect(source)
    assert result and len(commands) == 1
    assert source.read_bytes() == b"synthetic legacy document"
