from __future__ import annotations

import argparse
import hashlib
import html
import json
import subprocess
import sys
from collections import Counter
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from docx.opc.exceptions import PackageNotFoundError

from .collector import Placeholder, collect
from .output import validate_output_paths, write_new_outputs
from .review import review_findings


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="manuscript-placeholders",
        description="Collect manuscript placeholders into a revision checklist.",
    )
    parser.add_argument("manuscript")
    parser.add_argument(
        "--format", choices=("text", "markdown", "json", "html"), default="text"
    )
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--pattern",
        action="append",
        default=[],
        help="additional regular expression; repeat as needed",
    )
    parser.add_argument("--fail-on-findings", action="store_true")
    parser.add_argument("--docx-tables", action="store_true")
    parser.add_argument("--docx-comments", action="store_true")
    parser.add_argument("--docx-footnotes", action="store_true")
    parser.add_argument("--docx-headers", action="store_true")
    parser.add_argument("--docx-footers", action="store_true")
    parser.add_argument("--docx-endnotes", action="store_true")
    parser.add_argument("--docx-textboxes", action="store_true")
    parser.add_argument(
        "--revisions", choices=("current", "include-revisions"), default="current"
    )
    parser.add_argument(
        "--multiline",
        action="store_true",
        help="scan contiguous text lines within each section; markers are bounded to 8192 characters",
    )
    parser.add_argument(
        "--allow-cross-document",
        action="store_true",
        help="explicitly permit reviewed legacy or cross-document review files",
    )
    parser.add_argument(
        "--document-id",
        help="stable private identity for a manuscript moved to a new path",
    )
    parser.add_argument("--ledger", type=Path)
    parser.add_argument("--ledger-output", type=Path)
    parser.add_argument("--baseline", type=Path)
    args = parser.parse_args(argv)
    try:
        targets = [p for p in (args.output, args.ledger_output) if p]
        validate_output_paths(targets)
        coverage: dict[str, object] = {}
        items = collect(
            Path(args.manuscript),
            tuple(args.pattern),
            docx_tables=args.docx_tables,
            docx_comments=args.docx_comments,
            docx_footnotes=args.docx_footnotes,
            docx_headers=args.docx_headers,
            docx_footers=args.docx_footers,
            docx_endnotes=args.docx_endnotes,
            docx_textboxes=args.docx_textboxes,
            revisions=args.revisions,
            multiline=args.multiline,
            coverage=coverage,
            document_id=args.document_id,
        )
        source = Path(args.manuscript).expanduser().resolve()
        provenance = {
            "document_id": args.document_id
            or hashlib.sha256(str(source).encode()).hexdigest(),
            "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "scanner_version": "1.5.0",
            "coverage": coverage,
        }
        review = review_findings(
            [item.to_dict() for item in items],
            json.loads(args.ledger.read_text(encoding="utf-8")) if args.ledger else {},
            json.loads(args.baseline.read_text(encoding="utf-8"))
            if args.baseline
            else None,
            provenance=provenance,
            allow_cross_document=args.allow_cross_document,
        )
        review["coverage"] = coverage
        review["docx_structures"] = {
            name: "inapplicable"
            if source.suffix.casefold() != ".docx"
            else "included"
            if getattr(args, "docx_" + name)
            else "skipped unless explicitly enabled"
            for name in (
                "tables",
                "comments",
                "footnotes",
                "headers",
                "footers",
                "endnotes",
            )
        }
    except (
        subprocess.TimeoutExpired,
        UnicodeError,
        RuntimeError,
        OSError,
        ValueError,
        TypeError,
        ParseError,
        BadZipFile,
        PackageNotFoundError,
    ) as exc:
        print(f"manuscript-placeholders: {exc}", file=sys.stderr)
        return 2
    report = _render(items, args.format, Path(args.manuscript).name)
    report = report.replace(
        f"{len(items)} open items", f"{len(items)} observed items"
    ).replace(f"Open items: **{len(items)}**", f"Observed items: **{len(items)}**")
    if args.format == "json":
        payload = json.loads(report)
        payload["version"] = 2
        payload["provenance"] = provenance
        payload["review"] = review
        report = json.dumps(payload, indent=2, ensure_ascii=False) + "\n"
    elif args.format == "html":
        summaries = []
        for item in items:
            entry = review["ledger"]["items"][item.fingerprint]
            summaries.append(
                f"<li><strong>{html.escape(item.marker)}</strong>: {html.escape(entry['status'])} — {html.escape(entry.get('note', ''))}</li>"
            )
        comparison = review["comparison"]
        report = report.replace(
            "</main>",
            "<h2>Private resolution review</h2><p>"
            + html.escape(review["note"])
            + "</p><p>Status totals: "
            + html.escape(json.dumps(review["status_totals"]))
            + "</p><ul>"
            + "".join(summaries)
            + "</ul>"
            + f"<p>Compared with baseline: {len(comparison['new'])} new; {len(comparison['no_longer_observed'])} no longer observed; {len(comparison['unchanged'])} unchanged.</p>"
            + "<p>DOCX structure coverage: "
            + html.escape(
                "; ".join(
                    f"{name}: {state}"
                    for name, state in review["docx_structures"].items()
                )
            )
            + "</p><h2>Actual scan coverage</h2><pre>"
            + html.escape(json.dumps(coverage, indent=2, ensure_ascii=False))
            + "</pre><p>Unverified removals after coverage change: "
            + str(len(review["unverified_removals"]))
            + ". A missing marker is not proof of resolution.</p></main>",
        )
    else:
        report += "\nPrivate resolution review\n" + json.dumps(review, indent=2) + "\n"
    try:
        outputs = []
        if args.output:
            outputs.append((args.output, report))
        if args.ledger_output:
            outputs.append(
                (args.ledger_output, json.dumps(review["ledger"], indent=2) + "\n")
            )

        def deliver() -> None:
            sys.stdout.write(report)
            sys.stdout.flush()

        write_new_outputs(outputs, deliver=None if args.output else deliver)
        if args.output:
            print(f"Wrote {args.format} report to {args.output}")
    except OSError as exc:
        print(f"manuscript-placeholders: {exc}", file=sys.stderr)
        return 2
    return 1 if args.fail_on_findings and items else 0


def _render(items: tuple[Placeholder, ...], format_name: str, source: str) -> str:
    if format_name == "json":
        return (
            json.dumps(
                {
                    "source": source,
                    "count": len(items),
                    "placeholders": [item.to_dict() for item in items],
                },
                indent=2,
                ensure_ascii=False,
            )
            + "\n"
        )
    categories = Counter(item.category for item in items)
    if format_name == "markdown":
        lines = [
            "# Manuscript placeholder report",
            "",
            f"Source: `{source}`  ",
            f"Open items: **{len(items)}**",
            "",
        ]
        for category in sorted(categories):
            lines.extend((f"## {category} ({categories[category]})", ""))
            for item in (entry for entry in items if entry.category == category):
                lines.append(
                    f"- [ ] **{item.section} · {item.location} chars {item.start}:{item.end}:** `{item.marker}` — {item.context}"
                )
            lines.append("")
        return "\n".join(lines)
    if format_name == "html":
        cards = (
            "".join(
                f'<article><p class="category">{html.escape(item.category)}</p><h2>{html.escape(item.marker)}</h2><p>{html.escape(item.context)}</p><small>{html.escape(item.section)} · {html.escape(item.location)} chars {item.start}:{item.end}</small></article>'
                for item in items
            )
            or "<p>No placeholders found.</p>"
        )
        return f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Manuscript placeholders</title><style>body{{font:16px/1.55 Georgia,serif;background:#f3eee5;color:#29241f;margin:0}}main{{max-width:900px;margin:auto;padding:3rem 1rem}}article{{background:white;padding:1.25rem;margin:1rem 0;border-left:5px solid #9c552f}}.category,small{{color:#775d4d;font:13px system-ui}}h2{{font-size:1.15rem}}</style></head><body><main><h1>Manuscript placeholder report</h1><p>{html.escape(source)} · {len(items)} open items</p>{cards}</main></body></html>'
    lines = [f"Manuscript placeholders: {len(items)}", f"Source: {source}", ""]
    for item in items:
        lines.extend(
            (
                f"{item.category.upper()} · {item.section} · block {item.block} · {item.location} chars {item.start}:{item.end}",
                f"  {item.marker}",
                f"  {item.context}",
                "",
            )
        )
    return "\n".join(lines)


if __name__ == "__main__":
    raise SystemExit(main())
